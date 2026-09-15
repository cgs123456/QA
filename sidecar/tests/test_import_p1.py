"""P1 格式导入单测：dry-run → 确认 → 编译后端闭环。

- 真实样例：xlsx 经 openpyxl 实写、PDF 为手写最小文本页（pypdf 可读），
  均走 TestClient preview→commit 全程（DoD）。
- 拒绝路径：坏文件/超大格/空导入/扫描 PDF（422 明确提示，不静默跳过）。
- 恶意 Excel：公式格（data_only 无缓存→空）与注入前导（加 ' 中和+计数）。
- commit 篡改：未知角色/越界列/未知表/空映射/PDF 非法映射一律 400。
"""

import base64
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import openpyxl

from knowledge.parsers import excel_parser, pdf_parser

TEST_TOKEN = "test-token-" + "x" * 32


def _h():
    return {"Authorization": f"Bearer {TEST_TOKEN}"}


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def make_workbook(sheets: dict) -> bytes:
    """sheets: {表名: [[行...]]} → 真实 .xlsx 字节。"""
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for r in rows:
            ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


QA_SHEET = [
    ["问题", "答案", "分类"],
    ["发货周期是多久？", "三天内发出", "售后"],
    ["支持哪些支付？", "微信和支付宝", "售后"],
    ["退货期限是几天？", "七天内可退", "售后"],
    ["客服电话是多少？", "400-123-4567", "售后"],
    ["发票怎么开？", "备注抬头开电子票", "售后"],
]

FIELD_SHEET = [
    ["实体", "字段名", "字段值"],
    ["公司信息", "客服电话", "400-123-4567"],
]


def build_pdf(pages_lines: list) -> bytes:
    """最小单栏文本 PDF（ASCII 行；pypdf 可提取）。"""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    page_nums = list(range(4, 4 + len(pages_lines)))
    kids = " ".join(f"{n} 0 R" for n in page_nums)
    objects.append(
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages_lines)} >>".encode("ascii"))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    stream_nums = list(range(4 + len(pages_lines), 4 + 2 * len(pages_lines)))
    for pnum, snum in zip(page_nums, stream_nums):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {snum} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
            .encode("ascii"))
    for lines in pages_lines:
        ops, y = [], 700
        for ln in lines:
            esc = ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"BT /F1 12 Tf 50 {y} Td ({esc}) Tj ET")
            y -= 20
        stream = "\n".join(ops).encode("ascii")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream
                       + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n" % (len(objects) + 1))
    out.write(b"0000000000 65535 f \n")
    for o in offsets:
        out.write(b"%010d 00000 n \n" % o)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF"
              % (len(objects) + 1, xref))
    return out.getvalue()


TEXT_PAGE = [
    "1 Info",
    "Phone: 400-123-4567",
    "Address: Beijing Chaoyang district building 5 room 501 zip 100200",
    "Q: What is the shipping time?",
    "A: Ships within three days after payment confirmation received ok.",
]


def scanned_pdf() -> bytes:
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(612, 792)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


# ---------- Excel 解析/提案 ----------

def test_excel_preview_proposals_and_samples():
    data = make_workbook({"QA": QA_SHEET, "Fields": FIELD_SHEET})
    pv = excel_parser.parse_excel_preview(data, "s.xlsx")
    assert pv["format"] == "excel"
    by_sheet = {s["sheet"]: s for s in pv["sheets"]}
    qa = by_sheet["QA"]
    assert qa["header_row"] == 1
    assert qa["suggested_mapping"] == {"0": "standard_question",
                                       "1": "official_answer",
                                       "2": "category"}
    # 样例只给 3 行（5 行数据）。
    assert len(qa["columns"][0]["samples"]) == 3
    assert qa["columns"][0]["candidates"][0]["role"] == "standard_question"
    assert "置信" in qa["columns"][0]["candidates"][0]["reason"] or \
        "同义词" in qa["columns"][0]["candidates"][0]["reason"]
    fields = by_sheet["Fields"]
    assert fields["suggested_mapping"] == {"0": "entity", "1": "field_name",
                                           "2": "field_value"}


def test_excel_multi_sheet_independent_english_headers():
    data = make_workbook({
        "QA": QA_SHEET,
        "EN": [["standard_question", "official_answer"],
               ["How long?", "Three days of waiting time here ok"],
               ["Where?", "Beijing headquarter building number five"]],
    })
    pv = excel_parser.parse_excel_preview(data)
    by_sheet = {s["sheet"]: s for s in pv["sheets"]}
    # 角色名精确命中（0.95），与中文表独立提案。
    assert by_sheet["EN"]["suggested_mapping"] == {"0": "standard_question",
                                                   "1": "official_answer"}
    assert by_sheet["EN"]["columns"][0]["candidates"][0]["confidence"] == 0.95
    assert by_sheet["QA"]["suggested_mapping"]["0"] == "standard_question"


def test_excel_no_header_warns_and_manual_mapping_commits():
    data = make_workbook({"S": [["400-1111"], ["400-2222"], ["400-3333"]]})
    pv = excel_parser.parse_excel_preview(data)
    sheet = pv["sheets"][0]
    assert sheet["header_row"] is None
    assert any("表头" in w for w in sheet["warnings"])
    # 无 field_name 列 → 产不出字段（entity+field_value 不够），拒绝空导入。
    with pytest.raises(ValueError) as exc:
        excel_parser.apply_excel_mapping(data, {
            "S": {"0": "field_value", "__defaults__": {"entity": "公司信息"}}})
    assert "空导入" in str(exc.value)
    with pytest.raises(ValueError):
        excel_parser.apply_excel_mapping(data, {
            "S": {"0": "ignore", "__defaults__": {"entity": "X"}}})


def test_excel_sheet_defaults_entity():
    """整表同属一实体：__defaults__ 补 entity（有表头表的常规形态）。"""
    data = make_workbook({"S": [["字段名", "字段值"],
                                ["客服电话", "400-1111"],
                                ["售后电话", "400-2222"]]})
    out = excel_parser.apply_excel_mapping(data, {
        "S": {"0": "field_name", "1": "field_value",
              "__defaults__": {"entity": "公司信息"}}})
    assert len(out["field_items"]) == 2
    assert out["field_items"][0]["entity"] == "公司信息"
    with pytest.raises(ValueError):
        excel_parser.apply_excel_mapping(data, {
            "S": {"0": "field_name", "1": "field_value",
                  "__defaults__": {"entity": "  "}}})


def test_excel_formula_and_injection_protection():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QA"
    ws.append(["问题", "答案"])
    ws.append(["=1+1", "三天内发出"])  # 公式：data_only 下无缓存→空
    ws.append(["+12345", "微信和支付宝"])  # 注入前导：中和
    ws.append(["\uFF1Dcalc", "正常答案"])  # 全角＝：中和
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    pv = excel_parser.parse_excel_preview(data)
    assert any("中和" in w for w in pv["sheets"][0]["warnings"])
    out = excel_parser.apply_excel_mapping(
        data, {"QA": {"0": "standard_question", "1": "official_answer"}})
    questions = [q["standard_question"] for q in out["qa_items"]]
    # 公式格视为空（该行跳过），注入格加 ' 前缀入库。
    assert not any("1+1" in q for q in questions)
    assert "'+12345" in questions
    assert "'\uFF1Dcalc" in questions
    assert out["neutralized"] == 2


def test_excel_oversize_cell_rejected():
    data = make_workbook({"S": [["问题", "答案"], ["x" * 9001, "y"]]})
    with pytest.raises(ValueError) as exc:
        excel_parser.parse_excel_preview(data)
    assert "超大" in str(exc.value)


def test_excel_bad_file_rejected():
    with pytest.raises(ValueError):
        excel_parser.parse_excel_preview(b"this is not a zip")


# ---------- Excel preview→commit 全程 ----------

def test_excel_preview_commit_walk(api_client):
    client, _ = api_client
    data = make_workbook({"QA": QA_SHEET, "Fields": FIELD_SHEET})
    r = client.post("/knowledge/import/preview",
                    json={"format": "excel", "content_b64": _b64(data),
                          "filename": "qa.xlsx"},
                    headers=_h())
    assert r.status_code == 200, r.text
    body = r.json()
    assert len(body["sheets"]) == 2
    mapping = {s["sheet"]: s["suggested_mapping"] for s in body["sheets"]}

    r = client.post("/knowledge/import/commit",
                    json={"store_name": "P1库", "format": "excel",
                          "content_b64": _b64(data), "filename": "qa.xlsx",
                          "mapping": mapping},
                    headers=_h())
    assert r.status_code == 200, r.text
    stats = r.json()["stats"]
    assert stats["qa_inserted"] == 5
    assert stats["fields_upserted"] == 1

    _, conn = api_client
    n_qa = conn.execute("SELECT COUNT(*) FROM qa_pairs").fetchone()[0]
    n_f = conn.execute("SELECT COUNT(*) FROM fields").fetchone()[0]
    assert (n_qa, n_f) == (5, 1)
    row = conn.execute(
        "SELECT official_answer FROM qa_pairs WHERE standard_question='发票怎么开？'"
    ).fetchone()
    assert row[0] == "备注抬头开电子票"
    # preview 不写库：commit 前库为空由事务性保证（本用例 commit 一次即 5 条，
    # 无 preview 残留；重复 commit 同题则走 update/skip）。
    r2 = client.post("/knowledge/import/commit",
                     json={"store_name": "P1库2", "format": "excel",
                           "content_b64": _b64(data), "mapping": mapping},
                     headers=_h())
    assert r2.status_code == 200


def test_excel_commit_requires_mapping(api_client):
    client, _ = api_client
    data = make_workbook({"QA": QA_SHEET})
    r = client.post("/knowledge/import/commit",
                    json={"store_name": "X", "format": "excel",
                          "content_b64": _b64(data)},
                    headers=_h())
    assert r.status_code == 400


@pytest.mark.parametrize("mapping, hint", [
    ({"QA": {"0": "admin"}}, "未知"),
    ({"QA": {"99": "standard_question"}}, "越界"),
    ({"Nope": {"0": "standard_question"}}, "未知工作表"),
    ({"QA": "not-a-dict"}, "须为对象"),
    ({"QA": {"0": "ignore"}}, "空导入"),
    ({"QA": {"0": "standard_question", "1": "standard_question",
              "2": "official_answer"}}, "重复映射"),
    ({"QA": {"x": "standard_question"}}, "非整数"),
    ({"QA": {"0": "standard_question", "__defaults__": {"foo": "x"}}}, "不支持"),
])
def test_excel_commit_tampered_mapping_rejected(api_client, mapping, hint):
    client, _ = api_client
    data = make_workbook({"QA": QA_SHEET})
    r = client.post("/knowledge/import/commit",
                    json={"store_name": "X", "format": "excel",
                          "content_b64": _b64(data), "mapping": mapping},
                    headers=_h())
    assert r.status_code == 400, r.text
    assert hint in r.text


def test_excel_commit_empty_sheet_rejected(api_client):
    client, _ = api_client
    wb = openpyxl.Workbook()
    wb.active.title = "Empty"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/knowledge/import/commit",
                    json={"store_name": "X", "format": "excel",
                          "content_b64": _b64(buf.getvalue()),
                          "mapping": {"Empty": {"0": "standard_question"}}},
                    headers=_h())
    assert r.status_code == 400


def test_import_bad_base64_and_garbage(api_client):
    client, _ = api_client
    for fmt in ("excel", "pdf"):
        r = client.post("/knowledge/import/preview",
                        json={"format": fmt, "content_b64": "!!!not-b64!!!"},
                        headers=_h())
        assert r.status_code == 400
        r = client.post("/knowledge/import/commit",
                        json={"store_name": "X", "format": fmt,
                              "content_b64": _b64(b"junk-bytes-here")},
                        headers=_h())
        assert r.status_code == 400


# ---------- PDF 解析/预览 ----------

def test_pdf_preview_extracts_fields_and_qa():
    data = build_pdf([TEXT_PAGE])
    pv = pdf_parser.parse_pdf_preview(data, "info.pdf")
    assert pv["supported"] is True
    assert pv["n_pages"] == 1
    assert pv["unsupported_pages"] == []
    assert pv["layout_support"] == "single-column-only"
    assert pv["field_count"] == 2
    assert pv["qa_count"] == 1
    assert pv["field_samples"][0]["field_name"] == "Phone"
    assert pv["qa_samples"][0]["standard_question"] == "What is the shipping time?"
    # 章节标题 "1 Info" 成为 entity。
    assert pv["field_samples"][0]["entity"] == "1 Info"


def test_pdf_colon_rule_shape():
    from knowledge.parsers.pdf_parser import _COLON_RE

    assert _COLON_RE.match("Phone: 123").groups() == ("Phone", "123")
    assert _COLON_RE.match("客服电话：400").groups() == ("客服电话", "400")
    assert _COLON_RE.match("no colon here") is None
    assert _COLON_RE.match("x" * 21 + ": y") is None  # 名称限 20 字
    assert _COLON_RE.match(": novalue") is None
    assert _COLON_RE.match("客服电话：400") is not None


def test_pdf_scanned_rejected_explicitly():
    data = scanned_pdf()
    with pytest.raises(pdf_parser.UnsupportedPDFError) as exc:
        pdf_parser.parse_pdf_preview(data, "scan.pdf")
    assert "扫描" in str(exc.value)
    assert exc.value.pages == [0]


def test_pdf_orphan_q_before_heading_starts_new_entity():
    """裸 Q 后直接换标题：孤 Q 作废计数，标题正常切换域（不粘连）。"""
    from knowledge.parsers.pdf_parser import _extract_items

    pages = [{"index": 0, "chars": 100, "supported": True,
              "text": "Q: orphan question here\n2 Next\nPhone: 12345678901234567890123"}]
    out = _extract_items(pages, "doc")
    assert out["qa_items"] == []
    assert out["field_items"] == [{"entity": "2 Next", "field_name": "Phone",
                                   "field_value": "12345678901234567890123"}]
    assert out["skipped"] == 1


def test_pdf_mixed_pages_marked_not_silently_skipped():
    data = build_pdf([TEXT_PAGE, []])
    pv = pdf_parser.parse_pdf_preview(data, "mix.pdf")
    assert pv["supported"] is True
    assert pv["unsupported_pages"] == [1]
    assert pv["field_count"] == 2  # 文本页照常提取
    assert any("图片" in w or "扫描" in w for w in pv["warnings"])


# ---------- PDF preview→commit 全程 ----------

def test_pdf_preview_commit_walk(api_client):
    client, _ = api_client
    data = build_pdf([TEXT_PAGE])
    r = client.post("/knowledge/import/preview",
                    json={"format": "pdf", "content_b64": _b64(data),
                          "filename": "info.pdf"},
                    headers=_h())
    assert r.status_code == 200, r.text
    assert r.json()["qa_count"] == 1

    r = client.post("/knowledge/import/commit",
                    json={"store_name": "PDF库", "format": "pdf",
                          "content_b64": _b64(data), "filename": "info.pdf"},
                    headers=_h())
    assert r.status_code == 200, r.text
    stats = r.json()["stats"]
    assert stats["qa_inserted"] == 1
    assert stats["fields_upserted"] == 2
    assert stats["pdf_unsupported_pages"] == []

    _, conn = api_client
    row = conn.execute(
        "SELECT field_value FROM fields WHERE field_name='Phone'").fetchone()
    assert row[0] == "400-123-4567"


def test_pdf_scanned_preview_and_commit_422(api_client):
    client, _ = api_client
    b64 = _b64(scanned_pdf())
    r = client.post("/knowledge/import/preview",
                    json={"format": "pdf", "content_b64": b64,
                          "filename": "scan.pdf"},
                    headers=_h())
    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert detail["error"] == "unsupported"
    assert "扫描" in detail["message"]

    r = client.post("/knowledge/import/commit",
                    json={"store_name": "X", "format": "pdf",
                          "content_b64": b64, "filename": "scan.pdf"},
                    headers=_h())
    assert r.status_code == 422
    assert r.json()["detail"]["error"] == "unsupported"


def test_pdf_mixed_commit_warns_and_compiles(api_client):
    client, _ = api_client
    data = build_pdf([TEXT_PAGE, []])
    r = client.post("/knowledge/import/commit",
                    json={"store_name": "M", "format": "pdf",
                          "content_b64": _b64(data), "filename": "m.pdf"},
                    headers=_h())
    assert r.status_code == 200, r.text
    assert r.json()["stats"]["pdf_unsupported_pages"] == [1]
    assert r.json()["stats"]["fields_upserted"] == 2


def test_pdf_commit_bad_mapping_rejected(api_client):
    client, _ = api_client
    data = build_pdf([TEXT_PAGE])
    for mapping in ({"entity": 123}, {"foo": "x"}, {"entity": ""},
                    {"entity": "y" * 65}):
        r = client.post("/knowledge/import/commit",
                        json={"store_name": "X", "format": "pdf",
                              "content_b64": _b64(data),
                              "mapping": mapping},
                        headers=_h())
        assert r.status_code == 400, (mapping, r.text)


def test_pdf_no_items_commit_rejected(api_client):
    client, _ = api_client
    data = build_pdf([["This page has plenty of plain text over fifty characters "
                       "long but no colon fields or qa markers at all here"]])
    r = client.post("/knowledge/import/preview",
                    json={"format": "pdf", "content_b64": _b64(data)},
                    headers=_h())
    assert r.status_code == 200
    assert r.json()["qa_count"] == 0 and r.json()["field_count"] == 0
    r = client.post("/knowledge/import/commit",
                    json={"store_name": "X", "format": "pdf",
                          "content_b64": _b64(data)},
                    headers=_h())
    assert r.status_code == 400
